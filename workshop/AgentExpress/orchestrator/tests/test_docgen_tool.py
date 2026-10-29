"""Track B docgen. The Document generator (`source: "docgen"`, app/tools/docgen/).

What has to hold:
  * each format renders, from a small template (PPTX, DOCX) and from none (the built-in
    layout, and always for PDF): fields, split runs, repeated rows, paragraphs and slides,
    an empty list removing what uses it, a slide copy keeping its picture;
  * the function saves the file with the run and under the output prefix, links to one
    copy, and refuses a format not configured, a template name that leaves its folder, and
    a call with nowhere to save to; its S3 calls fit the service model;
  * the runtime sets runId and agent on every call to it, whatever the model sent;
  * both validators refuse a bad config (tests/fixtures/validation_cases.json runs them
    both; a few direct cases here);
  * IAM on both IaC paths: the template object or folder, run outputs, the output prefix,
    its own logs, nothing else; nothing at all unless a tool names it;
  * its libraries are pinned, installed by the deploy for Lambda, and the BFF's test run
    reaches it.
"""
from __future__ import annotations

import importlib.util
import io
import json
import re
import shutil
import sys
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import SimpleNamespace

import botocore.session
import pytest
from botocore.exceptions import ClientError
from botocore.validate import validate_parameters

docx = pytest.importorskip("docx")
pptx = pytest.importorskip("pptx")
pytest.importorskip("reportlab")

ORCH = Path(__file__).resolve().parent.parent
TOOL_DIR = ORCH / "app" / "tools" / "docgen"
S3_MODEL = botocore.session.get_session().get_service_model("s3")
DATA = {"title": "Quarterly report", "subtitle": "Third quarter", "owner": "Operations",
        "summary": {"total": 3, "ok": True}, "tags": ["north", "south"],
        "rows": [{"name": "North", "value": 1}, {"name": "South", "value": 2}]}


def _load(name: str, path: Path):
    sys.path.insert(0, str(TOOL_DIR))
    try:
        spec = importlib.util.spec_from_file_location(name, path)
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        spec.loader.exec_module(module)
    finally:
        sys.path.remove(str(TOOL_DIR))
    return module


@pytest.fixture(scope="module")
def render():
    return _load("ax_docgen_render", TOOL_DIR / "render.py")


# --- small templates, built here -------------------------------------------------------------

def _png() -> bytes:
    from PIL import Image
    out = io.BytesIO()
    Image.new("RGB", (4, 4), (200, 30, 30)).save(out, "PNG")
    return out.getvalue()


def docx_template() -> bytes:
    d = docx.Document()
    p = d.add_paragraph("Report: ")
    p.add_run("{{ti")                      # a field Word split over runs
    p.add_run("tle}} by {{ owner }}")
    d.add_paragraph("- {{rows.name}} = {{rows.value}}")
    d.add_paragraph("Gone {{empty.x}}")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Name"
    t.cell(0, 1).text = "Value"
    t.cell(1, 0).text = "{{rows.name}}"
    t.cell(1, 1).text = "{{rows.value}}"
    d.add_paragraph("Missing [{{nope}}], nested {{summary.total}}, list {{tags}}")
    d.sections[0].header.paragraphs[0].text = "Header {{owner}}"
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def pptx_template() -> bytes:
    from pptx.util import Inches
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "{{title}}"
    tb = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    tb.cell(0, 0).text = "Name"
    tb.cell(1, 0).text = "{{rows.name}}"
    tb.cell(1, 1).text = "{{rows.value}}"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Region {{rows.name}}"
    s2.placeholders[1].text = "Value {{rows.value}} of {{title}}"
    s2.shapes.add_picture(io.BytesIO(_png()), Inches(7), Inches(1), Inches(1), Inches(1))
    s3 = prs.slides.add_slide(prs.slide_layouts[1])
    s3.shapes.title.text = "Gone {{empty.x}}"
    s4 = prs.slides.add_slide(prs.slide_layouts[0])
    s4.shapes.title.text = "End {{owner}}"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _slides(body: bytes) -> list[list]:
    prs = pptx.Presentation(io.BytesIO(body))
    return [[sh.text_frame.text if sh.has_text_frame else
             ([[c.text for c in r.cells] for r in sh.table.rows] if sh.has_table else sh.shape_type)
             for sh in slide.shapes] for slide in prs.slides]


# --- rendering ------------------------------------------------------------------------------

def test_a_docx_template_is_filled_rows_and_paragraphs_repeat_and_an_empty_list_removes(render):
    out = docx.Document(io.BytesIO(render.render("docx", {**DATA, "empty": []}, docx_template())))
    assert [p.text for p in out.paragraphs] == [
        "Report: Quarterly report by Operations", "- North = 1", "- South = 2",
        "Missing [], nested 3, list north, south"]
    assert [[c.text for c in r.cells] for r in out.tables[0].rows] == [["Name", "Value"], ["North", "1"], ["South", "2"]]
    assert out.sections[0].header.paragraphs[0].text == "Header Operations"


def test_a_pptx_template_repeats_slides_and_rows_and_a_copy_keeps_its_picture(render):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    body = render.render("pptx", {**DATA, "empty": []}, pptx_template())
    assert _slides(body) == [
        ["Quarterly report", [["Name", ""], ["North", "1"], ["South", "2"]]],
        ["Region North", "Value 1 of Quarterly report", MSO_SHAPE_TYPE.PICTURE],
        ["Region South", "Value 2 of Quarterly report", MSO_SHAPE_TYPE.PICTURE],
        ["End Operations", ""],
    ]
    prs = pptx.Presentation(io.BytesIO(body))
    # The copy's picture points at an image through the copy's own relationship.
    for slide in list(prs.slides)[1:3]:
        pic = next(sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
        assert pic.image.blob == _png() or pic.image.content_type == "image/png"
    assert len(prs.slides) == 4


@pytest.mark.parametrize("fmt, magic", [("pdf", b"%PDF"), ("docx", b"PK"), ("pptx", b"PK")])
def test_each_format_renders_in_the_built_in_layout(render, fmt, magic):
    assert render.render(fmt, DATA, None).startswith(magic)


def test_the_built_in_layout_has_the_title_then_each_field_in_order(render):
    d = docx.Document(io.BytesIO(render.render("docx", DATA, None)))
    texts = [p.text for p in d.paragraphs]
    assert texts[:3] == ["Quarterly report", "Third quarter", "Owner"]
    assert "Total: 3" in texts and "north" in texts
    assert [[c.text for c in r.cells] for r in d.tables[0].rows] == [["Name", "Value"], ["North", "1"], ["South", "2"]]
    slides = _slides(render.render("pptx", DATA, None))
    assert slides[0] == ["Quarterly report", "Third quarter"]
    assert [s[0] for s in slides[1:]] == ["Owner", "Summary", "Tags", "Rows"]
    title, _sub, parts = render.sections(DATA)
    assert title == "Quarterly report" and [p[1] for p in parts] == ["text", "pairs", "bullets", "table"]


def test_values_are_text_controls_are_dropped_and_lists_are_bounded(render):
    assert render.as_text(None) == "" and render.as_text(True) == "true" and render.as_text(2.5) == "2.5"
    assert render.as_text(["a", 1]) == "a, 1" and render.as_text({"a": 1}) == '{"a": 1}'
    assert render.as_text("a\x00b\x07c\td") == "abc\td"
    d = docx.Document()
    d.add_paragraph("{{rows.name}}")
    tpl = io.BytesIO()
    d.save(tpl)
    with pytest.raises(render.RenderError, match="at most 1000"):
        render.render("docx", {"rows": [{"name": "x"}] * 1001}, tpl.getvalue())


def test_a_file_that_is_not_a_template_is_a_clear_error(render):
    with pytest.raises(render.RenderError, match=r"not a readable \.docx"):
        render.render("docx", DATA, b"not a zip")
    with pytest.raises(render.RenderError, match=r"not a readable \.pptx"):
        render.render("pptx", DATA, b"not a zip")


# --- the function, with S3 faked -------------------------------------------------------------

def _valid(op: str, params: dict) -> None:
    validate_parameters(params, S3_MODEL.operation_model(op).input_shape)


_PRECONDITION = {"Error": {"Code": "PreconditionFailed", "Message": "At least one of the pre-conditions "
                                                                    "you specified did not hold"},
                 "ResponseMetadata": {"HTTPStatusCode": 412}}


class FakeS3:
    def __init__(self, objects=None):
        self.objects = dict(objects or {})
        #: (bucket, key) -> (Metadata, ETag, ContentType) of what put_object wrote.
        self.meta: dict = {}
        self.puts: list[dict] = []
        self.signed: list[dict] = []
        self.lock = threading.Lock()
        self.writes = 0

    def head_object(self, **kw):
        _valid("HeadObject", kw)
        at = (kw["Bucket"], kw["Key"])
        if at not in self.objects:
            raise ClientError({"Error": {"Code": "403", "Message": "Forbidden"}}, "HeadObject")
        meta, etag, ctype = self.meta.get(at, ({}, '"0"', "application/octet-stream"))
        return {"ContentLength": len(self.objects[at]), "Metadata": dict(meta), "ETag": etag, "ContentType": ctype}

    def get_object(self, **kw):
        _valid("GetObject", kw)
        return {"Body": io.BytesIO(self.objects[(kw["Bucket"], kw["Key"])])}

    def list_objects_v2(self, **kw):
        _valid("ListObjectsV2", kw)
        return {"Contents": [{"Key": k, "Size": len(v)} for (b, k), v in sorted(self.objects.items())
                             if b == kw["Bucket"] and k.startswith(kw.get("Prefix", ""))], "IsTruncated": False}

    def put_object(self, **kw):
        _valid("PutObject", kw)
        # S3's conditional writes, each one atomic step: with If-None-Match: * an existing key
        # is refused (412), so two writers of the same key never both succeed; with If-Match a
        # key whose ETag is no longer that one is refused (412).
        at = (kw["Bucket"], kw["Key"])
        with self.lock:
            if kw.get("IfNoneMatch") == "*" and at in self.objects:
                raise ClientError(_PRECONDITION, "PutObject")
            if "IfMatch" in kw and (at not in self.objects or self.meta.get(at, (None, None))[1] != kw["IfMatch"]):
                raise ClientError(_PRECONDITION, "PutObject")
            self.writes += 1
            self.puts.append(kw)
            self.objects[at] = kw["Body"]
            self.meta[at] = (dict(kw.get("Metadata") or {}), f'"{self.writes}"', kw.get("ContentType", ""))

    def generate_presigned_url(self, op, Params, ExpiresIn):
        _valid("GetObject", Params)
        self.signed.append({"op": op, **Params, "expires": ExpiresIn})
        return f"https://{Params['Bucket']}.s3.us-east-1.amazonaws.com/{Params['Key']}?X-Amz-Expires={ExpiresIn}"


@pytest.fixture()
def dg(monkeypatch):
    monkeypatch.setenv("AWS_REGION", "us-east-1")
    for k in ("TEMPLATE_S3_URI", "OUTPUT_PREFIX"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("FORMATS", "pdf,pptx,docx")
    monkeypatch.setenv("ASSETS_BUCKET", "agentcore-x-assets-123")
    before = sys.modules.pop("render", None)
    module = _load("ax_docgen_handler", TOOL_DIR / "handler.py")
    sys.modules.pop("render", None)
    if before is not None:
        sys.modules["render"] = before
    module._s3 = FakeS3()
    yield module
    sys.modules.pop("ax_docgen_handler", None)


def _context(tool="generate_document"):
    return SimpleNamespace(client_context=SimpleNamespace(custom={"bedrockAgentCoreToolName": f"docs___{tool}"}))


def test_a_file_is_saved_with_the_run_and_linked(dg):
    out = dg.lambda_handler({"format": "pdf", "data": DATA, "runId": "run-123456", "agent": "writer"}, _context())
    assert out["s3Uri"] == "s3://agentcore-x-assets-123/runs/run-123456/outputs/Quarterly report.pdf"
    assert out["runFile"] == "runs/run-123456/outputs/Quarterly report.pdf" and out["template"] == "built-in layout"
    assert out["url"].startswith("https://") and out["expiresInSeconds"] == 3600
    put = dg._s3.puts[0]
    assert put["ContentType"] == "application/pdf" and put["Metadata"] == {"agent": "writer"}
    assert put["Body"].startswith(b"%PDF") and out["bytes"] == len(put["Body"])


def test_with_an_output_prefix_it_writes_both_and_links_to_the_prefix_copy(dg, monkeypatch):
    monkeypatch.setenv("OUTPUT_PREFIX", "s3://my-outputs/reports/")
    out = dg.lambda_handler({"format": "docx", "data": DATA, "fileName": "../../etc/summary",
                             "runId": "run-123456", "agent": "writer"}, _context())
    keys = [(p["Bucket"], p["Key"]) for p in dg._s3.puts]
    assert keys == [("agentcore-x-assets-123", "runs/run-123456/outputs/summary.docx"),
                    ("my-outputs", "reports/run-123456/summary.docx")]
    assert out["s3Uri"] == "s3://my-outputs/reports/run-123456/summary.docx"
    assert dg._s3.signed[-1]["Bucket"] == "my-outputs"


def test_a_template_folder_gives_a_named_template_and_never_one_outside_it(dg, monkeypatch):
    monkeypatch.setenv("TEMPLATE_S3_URI", "s3://my-templates/forms/")
    dg._s3.objects[("my-templates", "forms/report.docx")] = docx_template()
    ok = dg.lambda_handler({"format": "docx", "data": DATA, "template": "report.docx", "runId": "run-123456"}, _context())
    assert ok["template"] == "s3://my-templates/forms/report.docx"
    assert docx.Document(io.BytesIO(dg._s3.puts[-1]["Body"])).paragraphs[0].text == "Report: Quarterly report by Operations"
    for bad in ("../secret.docx", "a//b.docx", "./x.docx", "a\\b.docx"):
        assert "is not a name under the template folder" in dg.lambda_handler(
            {"format": "docx", "data": DATA, "template": bad, "runId": "run-123456"}, _context())["error"]
    assert "is not a .pptx file" in dg.lambda_handler(
        {"format": "pptx", "data": DATA, "template": "report.docx", "runId": "run-123456"}, _context())["error"]
    assert "cannot read the template s3://my-templates/forms/other.docx" in dg.lambda_handler(
        {"format": "docx", "data": DATA, "template": "other.docx", "runId": "run-123456"}, _context())["error"]
    # No name: the built-in layout.
    assert dg.lambda_handler({"format": "docx", "data": DATA, "runId": "run-123456"}, _context())["template"] == "built-in layout"


def test_one_template_object_is_used_for_its_own_format_only_and_pdf_never_reads_one(dg, monkeypatch):
    monkeypatch.setenv("TEMPLATE_S3_URI", "s3://my-templates/deck.pptx")
    dg._s3.objects[("my-templates", "deck.pptx")] = pptx_template()
    assert dg.lambda_handler({"format": "pptx", "data": DATA, "runId": "run-123456"}, _context())["template"] \
        == "s3://my-templates/deck.pptx"
    for fmt in ("pdf", "docx"):
        assert dg.lambda_handler({"format": fmt, "data": DATA, "runId": "run-123456"}, _context())["template"] == "built-in layout"


def test_what_it_refuses(dg, monkeypatch):
    monkeypatch.setenv("FORMATS", "pdf")
    assert "format must be one of pdf, not 'docx'" in dg.lambda_handler({"format": "docx", "data": DATA, "runId": "run-123456"}, _context())["error"]
    # A call for no run (or an id that is not one), and no output prefix: nowhere to save it.
    for run in ("", "../other", "x"):
        assert "nowhere to save the file" in dg.lambda_handler({"format": "pdf", "data": DATA, "runId": run}, _context())["error"]
    assert dg.lambda_handler({}, _context("delete_everything")) == {
        "error": "unknown tool 'delete_everything'", "availableTools": ["generate_document"]}
    assert "over 1 MB" in dg.lambda_handler({"format": "pdf", "data": {"t": "x" * 1_000_001}, "runId": "run-123456"}, _context())["error"]
    assert dg._s3.puts == []


def test_data_may_be_json_text_or_plain_text(dg):
    dg.lambda_handler({"format": "docx", "data": json.dumps({"title": "From text"}), "runId": "run-123456"}, _context())
    assert dg._s3.puts[-1]["Key"].endswith("/From text.docx")
    dg.lambda_handler({"format": "docx", "data": "just words", "runId": "run-123456"}, _context())
    assert "just words" in [p.text for p in docx.Document(io.BytesIO(dg._s3.puts[-1]["Body"])).paragraphs]


def _parallel(dg, calls: list[tuple[dict, str]]) -> list[dict]:
    """Every call at the same moment (a barrier releases them together), as a parallel
    stage's agents call the generator; each (event, tool key)."""
    start = threading.Barrier(len(calls))

    def one(call):
        event, key = call
        start.wait()
        return dg.lambda_handler(event, SimpleNamespace(client_context=SimpleNamespace(
            custom={"bedrockAgentCoreToolName": f"{key}___generate_document"})))
    with ThreadPoolExecutor(len(calls)) as pool:
        return list(pool.map(one, calls))


@pytest.mark.parametrize("calls,names", [
    # Invoice processing: three generator tools in one parallel stage, no name and no title.
    ([({"format": "pdf", "data": {"supplier": "Example Supplies", "total": 120}, "runId": "run-inv001",
        "agent": agent, "attempt": f"1-{agent}"}, key)
      for agent, key in (("remittance", "remittanceDoc"), ("credit_note", "creditNoteDoc"), ("dunning", "dunningDoc"))],
     {"remittanceDoc.pdf", "creditNoteDoc.pdf", "dunningDoc.pdf"}),
    # Staff onboarding: one generator tool called three times at once by one attempt of a step.
    ([({"format": "docx", "data": {"items": ["laptop", "badge"]}, "runId": "run-hr0042", "agent": "welcome_pack",
        "attempt": "1-a1b2c3"}, "welcomePackDoc")] * 3,
     {"welcomePackDoc.docx", "welcomePackDoc-2.docx", "welcomePackDoc-3.docx"}),
])
def test_three_parallel_calls_with_no_name_each_keep_their_own_file(dg, calls, names):
    out = _parallel(dg, calls)
    assert all("error" not in o for o in out), out
    assert {o["fileName"] for o in out} == names
    run = calls[0][0]["runId"]
    kept = {k for (b, k) in dg._s3.objects if b == "agentcore-x-assets-123"}
    assert kept == {f"runs/{run}/outputs/{n}" for n in names}
    # Each call links to its own file, and every file is a whole document.
    assert len({o["s3Uri"] for o in out}) == 3 and all(o["runFile"].endswith(o["fileName"]) for o in out)
    assert all(p["IfNoneMatch"] == "*" for p in dg._s3.puts)


@pytest.mark.parametrize("asked,ext,data", [
    ("audit-findings", "pdf", {"title": "Audit findings", "findings": 4}),       # a research report
    ("alert-summary.pptx", "pptx", {"title": "Device alerts", "open": 2}),       # IoT alerts
])
def test_three_parallel_calls_with_the_same_name_never_overwrite_each_other(dg, monkeypatch, asked, ext, data):
    monkeypatch.setenv("OUTPUT_PREFIX", "s3://my-outputs/reports/")
    calls = [({"format": ext, "data": {**data, "part": i}, "fileName": asked, "runId": "run-par123",
               "agent": "writer", "attempt": "1-d4e5f6"}, "docs") for i in range(3)]
    out = _parallel(dg, calls)
    stem = asked.removesuffix(f".{ext}")
    want = {f"{stem}.{ext}", f"{stem}-2.{ext}", f"{stem}-3.{ext}"}
    # The name asked for is honoured for the first; the others get the next free ones.
    assert {o["fileName"] for o in out} == want
    run_files = {k: v for (b, k), v in dg._s3.objects.items() if b == "agentcore-x-assets-123"}
    prefix_files = {k: v for (b, k), v in dg._s3.objects.items() if b == "my-outputs"}
    assert set(run_files) == {f"runs/run-par123/outputs/{n}" for n in want}
    assert set(prefix_files) == {f"reports/run-par123/{n}" for n in want}
    # Three different documents, none lost: each copy under the prefix is its run file.
    assert len(set(run_files.values())) == 3
    assert sorted(run_files.values()) == sorted(prefix_files.values())


def test_a_free_explicit_name_is_kept_and_a_taken_one_moves_to_the_next(dg):
    first = dg.lambda_handler({"format": "pdf", "data": DATA, "fileName": "summary", "runId": "run-123456"}, _context())
    again = dg.lambda_handler({"format": "pdf", "data": DATA, "fileName": "summary", "runId": "run-123456"}, _context())
    other_run = dg.lambda_handler({"format": "pdf", "data": DATA, "fileName": "summary", "runId": "run-654321"}, _context())
    assert (first["fileName"], again["fileName"], other_run["fileName"]) == ("summary.pdf", "summary-2.pdf", "summary.pdf")
    # A long name stays within the run file limit with its counter.
    long = dg.lambda_handler({"format": "pdf", "data": DATA, "fileName": "x" * 300, "runId": "run-123456"}, _context())
    longer = dg.lambda_handler({"format": "pdf", "data": DATA, "fileName": "x" * 300, "runId": "run-123456"}, _context())
    assert len(long["fileName"]) <= dg.MAX_NAME and len(longer["fileName"]) <= dg.MAX_NAME
    assert longer["fileName"].endswith("-2.pdf")


# --- a step's new attempt (gate revise, rerun_from) replaces its own earlier file --------------

def _step(agent_id: str):
    from app.common.base import Agent
    a = Agent()
    a.id, a.name, a.agentcore, a.model, a.temperature, a.max_tokens = agent_id, agent_id, {}, None, 0, 1000
    a.top_p, a.stop_sequences, a.system_prompt = None, [], "P"
    return a


def _attempt_call(dg, agent, state: dict, run: str, key: str, args: dict) -> dict:
    """One call of the generator from one run of `agent`: the step's context sets the scope
    (app/common/context.py), the runtime sets runId, agent and attempt from it
    (gateway/client.py _run_args), and the Gateway passes the call to the function."""
    from app.common.context import AgentContext
    from app.features.gateway import client
    AgentContext(agent, {"topic": "t", **state}, {"configurable": {"thread_id": run}})
    event = client._run_args(SimpleNamespace(name=f"{key}___generate_document"), args)
    return dg.lambda_handler(event, SimpleNamespace(client_context=SimpleNamespace(
        custom={"bedrockAgentCoreToolName": f"{key}___generate_document"})))


def _delivered(dg, monkeypatch, run: str, target: dict) -> list[dict]:
    """What a deliver step's target sends: the run's files (app/common/deliver.py run_files,
    read from the same bucket) through the target's filters."""
    from app.common import deliver
    monkeypatch.setattr(deliver, "_client", lambda service, region="": dg._s3)
    monkeypatch.setattr(deliver.s3_presign, "client", lambda region="": dg._s3)
    return deliver.select(deliver.run_files(run), target)


@pytest.mark.parametrize("how,agent_id,key,args,first,revised,prefix", [
    # Invoice processing: the remittance advice is sent back at its gate (revise) and written again.
    ("revise", "remittance", "remittanceDoc", {"format": "pdf", "fileName": "remittance-advice"},
     {"title": "Remittance advice", "invoice": "INV-1001", "total": 120},
     {"title": "Remittance advice", "invoice": "INV-1001", "total": 118, "note": "early payment discount"}, ""),
    # HR onboarding: the welcome pack (no name, so the tool's key) is re-run from its step (rerun_from).
    ("rerun_from", "welcome_pack", "welcomePackDoc", {"format": "docx"},
     {"items": ["laptop", "badge"]}, {"items": ["laptop", "badge", "parking pass"]}, "s3://my-outputs/onboarding/"),
])
def test_a_steps_new_attempt_replaces_its_own_earlier_file_and_deliver_sends_only_that(
        dg, monkeypatch, how, agent_id, key, args, first, revised, prefix):
    if prefix:
        monkeypatch.setenv("OUTPUT_PREFIX", prefix)
    run, step = "run-attempt01", _step(agent_id)
    # Another step of the run already saved a file under the same name: never replaced.
    other = _attempt_call(dg, _step("auditor"), {}, run, key, {**args, "data": {"title": "Audit"}})
    one = _attempt_call(dg, step, {}, run, key, {**args, "data": first})
    draft = dg._s3.objects[("agentcore-x-assets-123", one["runFile"])]
    # The gate's revise, or a rerun_from: the step runs again with one more version in its
    # history and the reviewer's note as feedback.
    later = {"history": {agent_id: [{"version": 1, "output": "{}"}]}, "feedback": {agent_id: "Fix the total."}}
    two = _attempt_call(dg, step, later, run, key, {**args, "data": revised})
    assert "error" not in two, two
    assert two["fileName"] == one["fileName"] != other["fileName"], (how, one, two, other)
    run_files = {k: v for (b, k), v in dg._s3.objects.items() if b == "agentcore-x-assets-123"}
    assert set(run_files) == {one["runFile"], other["runFile"]}
    assert run_files[two["runFile"]] != draft and draft not in run_files.values()
    # The replacing write was conditional on the draft it replaced.
    assert dg._s3.puts[-2 if prefix else -1].get("IfMatch")
    if prefix:
        copies = {k: v for (b, k), v in dg._s3.objects.items() if b == "my-outputs"}
        assert copies[f"onboarding/{run}/{two['fileName']}"] == run_files[two["runFile"]]
    sent = _delivered(dg, monkeypatch, run, {"type": "s3", "uri": "s3://b-1/out/", "producedBy": [agent_id]})
    assert [f["name"] for f in sent] == [two["fileName"]]
    assert dg._s3.objects[("agentcore-x-assets-123", sent[0]["key"])] == run_files[two["runFile"]] != draft
    # A third attempt (rerun_from after the revise) replaces it again: still one file.
    third = {"history": {agent_id: [{"version": 1}, {"version": 2}]}}
    three = _attempt_call(dg, step, third, run, key, {**args, "data": {**revised, "v": 3}})
    assert three["fileName"] == one["fileName"]
    assert len(_delivered(dg, monkeypatch, run, {"type": "s3", "uri": "s3://b-1/out/", "producedBy": [agent_id]})) == 1


@pytest.mark.parametrize("asked,ext,data", [
    ("literature-review", "pdf", {"title": "Literature review", "papers": 12}),   # research report
    ("sensor-digest", "docx", {"title": "Sensor digest", "alerts": 5}),           # IoT alerts
])
def test_three_parallel_calls_of_a_new_attempt_replace_the_old_files_and_never_each_other(dg, asked, ext, data):
    # The earlier attempt of the step made two files of that name.
    for part in (1, 2):
        dg.lambda_handler({"format": ext, "data": {**data, "draft": part}, "fileName": asked, "runId": "run-par456",
                           "agent": "writer", "attempt": "1-old000"}, _context())
    old = {v for (b, _k), v in dg._s3.objects.items() if b == "agentcore-x-assets-123"}
    assert len(old) == 2
    calls = [({"format": ext, "data": {**data, "part": i}, "fileName": asked, "runId": "run-par456",
               "agent": "writer", "attempt": "2-new000"}, "docs") for i in range(3)]
    out = _parallel(dg, calls)
    assert all("error" not in o for o in out), out
    assert {o["fileName"] for o in out} == {f"{asked}.{ext}", f"{asked}-2.{ext}", f"{asked}-3.{ext}"}
    kept = {k: v for (b, k), v in dg._s3.objects.items() if b == "agentcore-x-assets-123"}
    assert len(kept) == 3 and len(set(kept.values())) == 3 and not old & set(kept.values())
    assert {dg._s3.meta[("agentcore-x-assets-123", k)][0]["attempt"] for k in kept} == {"2-new000"}


def test_a_file_of_another_step_with_no_attempt_or_unreadable_is_never_replaced(dg, monkeypatch):
    base = {"format": "pdf", "data": DATA, "fileName": "summary", "runId": "run-123456"}
    dg.lambda_handler({**base, "agent": "writer"}, _context())                          # no attempt recorded
    dg.lambda_handler({**base, "agent": "checker", "attempt": "1-aaa"}, _context())      # another step
    third = dg.lambda_handler({**base, "agent": "writer", "attempt": "2-bbb"}, _context())
    assert third["fileName"] == "summary-3.pdf"

    def denied(**kw):
        raise ClientError({"Error": {"Code": "403", "Message": "Forbidden"}}, "HeadObject")
    # summary-3 is the writer's own from an earlier attempt, but who wrote it cannot be read.
    monkeypatch.setattr(dg._s3, "head_object", denied)
    fourth = dg.lambda_handler({**base, "agent": "writer", "attempt": "3-ccc"}, _context())
    assert fourth["fileName"] == "summary-4.pdf"
    assert len([k for (b, k) in dg._s3.objects if b == "agentcore-x-assets-123"]) == 4
    assert not any(p.get("IfMatch") for p in dg._s3.puts)


def test_each_run_of_a_step_is_its_own_attempt():
    from app.common.context import AgentContext
    from app.features.observability import scope
    step, cfg = _step("writer"), {"configurable": {"thread_id": "run-abcdef"}}
    try:
        first = AgentContext(step, {"topic": "t"}, cfg).attempt
        assert scope.get_scope().attempt == first and first.startswith("1-")
        # The same version again (an attempt that failed adds no history) is still a new attempt.
        again = AgentContext(step, {"topic": "t"}, cfg).attempt
        revised = AgentContext(step, {"topic": "t", "history": {"writer": [{"version": 1}]}}, cfg).attempt
        assert len({first, again, revised}) == 3 and revised.startswith("2-")
    finally:
        scope.set_scope("", "")


def test_with_an_output_prefix_and_no_run_files_the_link_is_the_file_written(dg, monkeypatch):
    monkeypatch.delenv("ASSETS_BUCKET")
    monkeypatch.setenv("OUTPUT_PREFIX", "s3://my-outputs/reports/")
    out = dg.lambda_handler({"format": "pdf", "data": DATA, "fileName": "summary", "runId": "run-123456"}, _context())
    assert [(p["Bucket"], p["Key"]) for p in dg._s3.puts] == [("my-outputs", "reports/run-123456/summary.pdf")]
    assert out["s3Uri"] == "s3://my-outputs/reports/run-123456/summary.pdf" and out["runFile"] == ""
    assert dg._s3.signed[-1]["Key"] == "reports/run-123456/summary.pdf"
    # With no run, the file is directly under the prefix, as before.
    loose = dg.lambda_handler({"format": "pdf", "data": DATA, "fileName": "summary"}, _context())
    assert loose["s3Uri"] == "s3://my-outputs/reports/summary.pdf"



def test_the_safe_name_is_the_run_files_rule(dg):
    sys.path.insert(0, str(ORCH))
    from app.common import artifacts
    for name in ("../../x.pdf", "a b(1).pdf", "..", "", "x" * 300 + ".docx", "q?.pptx"):
        assert dg.safe_name(name) == artifacts.safe_name(name)


# --- the runtime sets runId and agent ----------------------------------------------------------

def test_the_runtime_sets_the_run_and_agent_on_every_call_to_it():
    from app.features.gateway import client
    from app.features.observability import scope
    scope.set_scope("run-abcdef", "writer")
    try:
        tool = SimpleNamespace(name="docs___generate_document")
        assert client._run_args(tool, {"format": "pdf", "runId": "someone-elses", "agent": "x"}) == {
            "format": "pdf", "runId": "run-abcdef", "agent": "writer", "attempt": ""}
        scope.set_scope("run-abcdef", "writer", attempt="2-abc123")
        assert client._run_args(tool, {"format": "pdf", "attempt": "forged"})["attempt"] == "2-abc123"
        other = SimpleNamespace(name="web___open_page")
        assert client._run_args(other, {"url": "u", "runId": "kept"}) == {"url": "u", "runId": "kept"}
    finally:
        scope.set_scope("", "")
    assert "args = _run_args(tool, args)" in (ORCH / "app" / "features" / "gateway" / "client.py").read_text()


# --- the validators ---------------------------------------------------------------------------

def _issues(tool: dict, files: bool = True) -> list[str]:
    import validate_build
    wf = json.loads((ORCH / "app" / "workflow.json").read_text())
    wf["tools"]["docs"] = tool
    if files:
        wf["orchestrator"]["files"] = {}
    if "${storage}" in json.dumps(tool):  # the token needs the build's storage (bff/storage_check.py)
        wf["orchestrator"]["storage"] = {"create": True}
    return [f"{i['path']}: {i['message']}" for i in validate_build.validate(wf) if i["path"].startswith("tools.docs")]


PRESET = {"type": "lambda", "description": "Docs", "source": "docgen", "call": "generate_document", "arg": "data",
          "formats": ["pdf", "pptx", "docx"],
          "toolSchema": [{"name": "generate_document", "properties": {
              "format": {"type": "string", "required": True}, "data": {"type": "object", "required": True},
              "runId": {"type": "string", "required": False}, "agent": {"type": "string", "required": False}}}]}


def test_the_validator_accepts_the_preset_and_each_s3_shape():
    assert _issues(PRESET) == []
    for t, o in (("s3://b-1/t.docx", "s3://b-1/"), ("s3://b-1/forms/", "s3://b-1/o/p/"),
                 ("${storage}/templates/", "${storage}/out/"), ("s3://b-1/", None)):
        assert _issues({**PRESET, "templateS3Uri": t, **({"outputPrefix": o} if o else {})}) == [], (t, o)


@pytest.mark.parametrize("change, fragment", [
    ({"templateS3Uri": "s3://b-1/t.pdf"}, "must be one template"),
    ({"templateS3Uri": "s3://b-1/t*/"}, "must be one template"),
    ({"templateS3Uri": "https://b-1/t.docx"}, "must be one template"),
    ({"outputPrefix": "s3://b-1/out"}, "ending in /"),
    ({"outputPrefix": "s3://B/"}, "ending in /"),
    ({"outputPrefix": "s3://agentcore-x-assets-1/"}, "framework's own buckets"),
    ({"formats": ["docx"], "templateS3Uri": "s3://b-1/t.pptx"}, 'must include "pptx"'),
    ({"formats": []}, "list at least one format"),
    ({"formats": ["pdf", "pdf"]}, "is listed twice"),
    ({"formats": ["html"]}, '"html" is not one of: pdf, pptx, docx'),
])
def test_the_validator_rejects_bad_configs(change, fragment):
    assert any(fragment in i for i in _issues({**PRESET, **change})), _issues({**PRESET, **change})


def test_the_validator_needs_run_files_and_formats_and_one_tool_name():
    assert _issues(PRESET, files=False) == [
        ("tools.docs.source: the Document generator saves each file with its run: set orchestrator.files "
         "(even {}) so this deployment keeps run files")]
    assert any("needs `formats`" in i for i in _issues({k: v for k, v in PRESET.items() if k != "formats"}))
    renamed = {**PRESET, "call": "render", "toolSchema": [{**PRESET["toolSchema"][0], "name": "render"}]}
    assert any('publishes one tool, "generate_document"' in i for i in _issues(renamed))
    props = PRESET["toolSchema"][0]["properties"]
    undeclared = {**PRESET, "toolSchema": [{**PRESET["toolSchema"][0],
                                            "properties": {k: v for k, v in props.items() if k not in ("runId", "agent")}}]}
    assert _issues(undeclared) == [
        ("tools.docs.toolSchema[0].properties: declare runId and agent (optional strings): the Gateway passes the "
         "function only the arguments its schema declares, and the workflow sets runId and agent on every call so "
         "each file is saved with its run")]


# --- IAM and resources, both IaC paths ----------------------------------------------------------

TF = (ORCH / "terraform" / "docgen_tool.tf").read_text()
TOOLS_TF = (ORCH / "terraform" / "tools.tf").read_text()
TS = (ORCH / "cdk" / "lib" / "docgen-tool.ts").read_text()
PLANE_TS = (ORCH / "cdk" / "lib" / "tool-plane.ts").read_text()
STACK_TS = (ORCH / "cdk" / "lib" / "orchestrator-stack.ts").read_text()


def _code(src: str) -> str:
    """Source without its comments: the checks are about what is granted, not prose."""
    src = re.sub(r"/\*[\s\S]*?\*/", "", src)
    return "\n".join(ln for ln in src.splitlines() if not ln.strip().startswith(("#", "//")))


def _tf_statements() -> dict:
    """Each statement in docgen_tool.tf's docgen_statements: Sid -> (actions, resource text)."""
    out = {}
    for block in TF.split("Sid      = ")[1:]:
        sid = block.split('"')[1]
        actions = set(re.findall(r'"((?:s3|logs):[A-Za-z]+)"', block.split("Resource")[0]))
        out[sid] = (actions, block.split("Resource =", 1)[1].split("\n", 1)[0].strip())
    return out


def test_terraform_grants_only_the_template_run_outputs_output_prefix_and_logs():
    st = _tf_statements()
    assert set(st) == {"Logs", "ReadItsTemplates", "SaveWithTheRun", "WriteItsOutputs"}
    assert st["Logs"] == ({"logs:CreateLogStream", "logs:PutLogEvents"}, '["${aws_cloudwatch_log_group.tool[n].arn}:*"]')
    assert st["ReadItsTemplates"] == ({"s3:GetObject"}, "[arn]")
    assert st["SaveWithTheRun"] == ({"s3:PutObject", "s3:GetObject"}, '["${try(aws_s3_bucket.assets[0].arn, "")}/runs/*/outputs/*"]')
    assert st["WriteItsOutputs"] == ({"s3:PutObject", "s3:GetObject"}, "[arn]")
    # One object, or everything under a folder; the output prefix is a folder.
    assert ('endswith(t.template_s3_uri, "/") ? "arn:aws:s3:::${trimprefix(t.template_s3_uri, "s3://")}*" : '
            '"arn:aws:s3:::${trimprefix(t.template_s3_uri, "s3://")}"') in TF
    assert '"arn:aws:s3:::${trimprefix(t.output_prefix, "s3://")}*" if t.output_prefix != ""' in TF
    # Read as well as write the run's files, with an output prefix too: who wrote a taken name.
    assert 'Action   = ["s3:PutObject", "s3:GetObject"]\n      Resource = ["${try(aws_s3_bucket.assets[0].arn, "")}/runs/*/outputs/*"]' in TF
    assert "s3:ListBucket" not in _code(TF) and 'Resource = "*"' not in TF and "s3:*" not in TF


def test_cdk_grants_the_same():
    for sid in ("ReadItsTemplates", "SaveWithTheRun", "WriteItsOutputs"):
        assert f'sid: "{sid}"' in TS
    assert 'actions: ["s3:GetObject"], resources: [s3ObjectArn(spec.templateS3Uri)]' in TS
    assert 'sid: "SaveWithTheRun",\n    actions: ["s3:PutObject", "s3:GetObject"],\n    resources: [runOutputs],' in TS
    assert 'actions: ["s3:PutObject", "s3:GetObject"], resources: [s3ObjectArn(spec.outputPrefix)]' in TS
    assert "const runOutputs = `${assetsBucketArn}/runs/*/outputs/*`;" in TS
    assert "logGroup.grantWrite(role)" in TS and "s3:ListBucket" not in _code(TS) and '"*"' not in _code(TS)


def test_both_paths_give_it_60_seconds_512_mb_and_the_same_environment():
    assert 'memory_size = each.value.source == "docgen" ? 512 : 256' in TOOLS_TF
    assert 'contains(["codeInterpreter", "browser", "docgen"], each.value.source) ? 60 : 30' in TOOLS_TF
    for var in ('FORMATS       = join(",", each.value.formats)', 'ASSETS_BUCKET = try(aws_s3_bucket.assets[0].id, "")',
                "TEMPLATE_S3_URI = each.value.template_s3_uri, OUTPUT_PREFIX = each.value.output_prefix"):
        assert var in TOOLS_TF
    assert "timeout: cdk.Duration.seconds(60)," in TS and "memorySize: 512," in TS
    for var in ('FORMATS: (spec.formats ?? []).join(",")', "ASSETS_BUCKET: assets.bucketName",
                "TEMPLATE_S3_URI: spec.templateS3Uri", "OUTPUT_PREFIX: spec.outputPrefix"):
        assert var in TS


def test_nothing_exists_unless_a_tool_names_it():
    assert re.search(r'builtin_lambda_docgen_tools\s*=\s*\{ for n, t in local\.builtin_lambda_tools : '
                     r'n => t if t\.source == "docgen" \}', TOOLS_TF)
    assert "for_each = local.builtin_lambda_docgen_tools" in TF
    assert "count = length(local.builtin_lambda_docgen_tools) > 0 ? 1 : 0" in TF
    for loc in ("docgen_template_arn", "docgen_output_arn", "docgen_statements"):
        assert re.search(rf"{loc}\s*=\s*\{{ for n, t in local\.builtin_lambda_docgen_tools", TF), loc
    # The other functions keep their memory, and pricing keeps exactly its environment.
    assert "PRICING_API_REGION = \"us-east-1\"" in TOOLS_TF
    assert "if (spec.source === DOCGEN_SOURCE) {" in PLANE_TS
    # The run files are passed to the tool plane only when the deployment keeps them.
    assert "...(assetsBucket && keepDays !== null" in STACK_TS


def test_both_paths_publish_the_attempt_the_runtime_sets_without_the_config_declaring_it():
    # CDK: generate_document gets `attempt` unless the config declares it; nothing else changes.
    assert "spec.source === DOCGEN_SOURCE ? withDocgenArgs(spec.toolSchema) : spec.toolSchema" in PLANE_TS
    assert 'attempt: { type: "string", required: false, description: "Set by the workflow on every call; leave it out." }' in TS
    assert "s.name !== DOCGEN_TOOL ? s :" in TS and "!(k in (s.properties ?? {}))" in TS
    # Terraform: the same argument, on the same condition (planned in the test below).
    assert 'attempt = { type = "string", description = "Set by the workflow on every call; leave it out." }' in TF
    assert ('if try(t.source, "") == "docgen" && try(s.name, "") == "generate_document" '
            '&& !contains(keys(try(s.properties, {})), pname)]') in TOOLS_TF
    # And the runtime's argument is the one both publish.
    assert '"attempt": s.attempt' in (ORCH / "app" / "features" / "gateway" / "client.py").read_text()


_TF_SCHEMA_RUN = '''
run "docgen_publishes_the_attempt" {
  command = plan
  plan_options {
    target = [aws_iam_role_policy.tool_lambda_docgen]
  }
  assert {
    condition     = [for p in local.lambda_tools["reportDoc"].tool_schema[0].properties : p.name if p.name == "attempt"] == ["attempt"]
    error_message = "The Document generator's schema does not publish attempt once."
  }
  assert {
    condition     = [for p in local.lambda_tools["declaredDoc"].tool_schema[0].properties : p.description if p.name == "attempt"] == ["mine"]
    error_message = "An attempt the config declares was replaced or published twice."
  }
  assert {
    condition     = length([for p in local.lambda_tools["notDocgen"].tool_schema[0].properties : p if p.name == "attempt"]) == 0
    error_message = "A Lambda tool that is not the Document generator got the attempt argument."
  }
  assert {
    condition     = contains([for s in jsondecode(aws_iam_role_policy.tool_lambda_docgen["reportDoc"].policy).Statement : s.Action if s.Sid == "SaveWithTheRun"][0], "s3:GetObject")
    error_message = "With an output prefix the Document generator cannot read who wrote a run file."
  }
}
'''


@pytest.mark.skipif(not shutil.which("terraform") or not (ORCH / "terraform" / ".terraform").exists(),
                    reason="terraform, initialised, is needed for the offline plan")
def test_terraform_plans_the_published_attempt_and_the_run_file_read(tmp_path):
    from test_v139_wiring import _shipped_with_v139, _tf_copy
    wf = _shipped_with_v139(True)
    lam = {"type": "lambda", "description": "Own function", "lambdaArn": "arn:aws:lambda:us-east-1:123456789012:function:own",
           "toolSchema": [{"name": "lookup", "properties": {"q": {"type": "string", "required": True}}}]}
    wf["tools"]["notDocgen"] = lam
    declared = json.loads(json.dumps(wf["tools"]["reportDoc"]))
    declared["toolSchema"][0]["properties"]["attempt"] = {"type": "string", "description": "mine"}
    wf["tools"]["declaredDoc"] = declared
    fixture = (ORCH / "tests" / "fixtures" / "v139_wiring.tftest.hcl").read_text()
    r = _tf_copy(tmp_path, wf, fixture.split('run "', 1)[0] + _TF_SCHEMA_RUN)
    assert r.returncode == 0, r.stdout[-6000:] + r.stderr[-2000:]
    assert re.search(r"Success! 1 passed, 0 failed", r.stdout), r.stdout[-2000:]


def test_both_paths_refuse_no_run_files_missing_libraries_and_an_unresolved_storage_token():
    assert "condition     = local.files_on" in TF
    assert 'fileexists("${path.module}/../app/tools/docgen/pptx/__init__.py")' in TF
    assert '!startswith(t.template_s3_uri, "$${storage}")' in TF
    assert "set orchestrator.files (even {})" in TS and 'path.join(folder, "pptx", "__init__.py")' in TS
    assert "needs the build's storage (orchestrator.storage)" in TS


# --- libraries, the deploy, the test run --------------------------------------------------------

#: Each pinned library and its licence (PyPI metadata, checked Oct 2026): all permissive.
LICENCES = {"python-pptx": "MIT", "python-docx": "MIT", "reportlab": "BSD", "lxml": "BSD-3-Clause",
            "pillow": "MIT-CMU", "XlsxWriter": "BSD-2-Clause", "typing_extensions": "PSF-2.0",
            "charset-normalizer": "MIT"}


def test_every_library_is_pinned_and_licensed_permissively():
    lines = [ln.strip() for ln in (TOOL_DIR / "requirements.txt").read_text().splitlines()
             if ln.strip() and not ln.startswith("#")]
    assert all(re.fullmatch(r"[A-Za-z0-9_.-]+==[0-9][0-9A-Za-z.]*", ln) for ln in lines), lines
    assert {ln.split("==")[0] for ln in lines} == set(LICENCES)


def test_the_deploy_installs_its_libraries_for_lambdas_python_once():
    sys.path.insert(0, str(ORCH / "deployer"))
    try:
        import runner
    finally:
        sys.path.pop(0)
    seen = []
    real_run = runner.run
    runner.run = lambda cmd, cwd, env=None, keep=60: seen.append(cmd)
    try:
        wf = {"tools": {"a": {"type": "lambda", "source": "docgen"}, "b": {"type": "lambda", "source": "docgen"},
                        "c": {"type": "lambda", "source": "browser"}, "d": {"type": "mcp", "source": "docgen"}}}
        assert runner.install_builtin_requirements(wf) == ["docgen"]
        assert runner.install_builtin_requirements({"tools": {"p": {"type": "lambda", "source": "pricing"}}}) == []
    finally:
        runner.run = real_run
    cmd = seen[0]
    assert len(seen) == 1 and cmd[1:4] == ["-m", "pip", "install"]
    assert cmd[cmd.index("-r") + 1].endswith("app/tools/docgen/requirements.txt")
    assert cmd[cmd.index("-t") + 1].endswith("app/tools/docgen")
    assert cmd[cmd.index("--platform") + 1] == "manylinux2014_x86_64" and "--only-binary=:all:" in cmd
    assert "install_builtin_requirements(bundle.get(\"workflow\") or {})" in (ORCH / "deployer" / "runner.py").read_text()


def test_the_function_imports_only_the_runtime_its_folder_and_its_libraries():
    import ast
    mods = set()
    for f in ("handler.py", "render.py"):
        for node in ast.walk(ast.parse((TOOL_DIR / f).read_text())):
            if isinstance(node, ast.Import):
                mods |= {a.name.split(".")[0] for a in node.names}
            elif isinstance(node, ast.ImportFrom) and node.module:
                mods.add(node.module.split(".")[0])
    assert mods - set(sys.stdlib_module_names) == {"boto3", "botocore", "render", "docx", "pptx", "reportlab"}


def test_the_source_is_in_the_vocabulary_and_the_shipped_sample_does_not_use_it():
    vocab = json.loads((ORCH / "app" / "vocabulary.json").read_text())
    assert "docgen" in vocab["builtinLambdaSource"]["values"]
    assert vocab["docgenFormats"]["values"] == ["pdf", "pptx", "docx"]
    workflow = json.loads((ORCH / "app" / "workflow.json").read_text())
    assert "docgen" not in {t.get("source") for t in workflow["tools"].values()}
    assert not any(k in t for t in workflow["tools"].values() for k in ("formats", "templateS3Uri", "outputPrefix"))


moto = pytest.importorskip("moto")
from test_builds import call, mark_deployed, project  # noqa: E402 - after the moto importorskip
from test_builds import env as env  # noqa: E402 - after the moto importorskip


def test_the_test_run_calls_the_deployed_document_generator(env, monkeypatch):
    import codecheck
    p = project(tools={"docs": PRESET})
    p["workflow"].setdefault("orchestrator", {})["files"] = {}
    assert call(env, "PUT /api/builds/{id}", params={"id": "pclaims01"}, body={"project": p})[0] == 200
    call(env, "POST /api/builds/{id}/deploy", params={"id": "pclaims01"}, body={"tool": "cdk"})
    agent = mark_deployed(env, "pclaims01")
    seen = {}

    def fake_invoke(client, fn, key, tool, event):
        seen.update(fn=fn, key=key, tool=tool, event=event)
        return {"ok": True, "status": 200, "error": "", "output": "{}", "log": "", "ms": 3}
    monkeypatch.setattr(codecheck, "invoke", fake_invoke)
    event = {"format": "pdf", "data": {"title": "t"}, "runId": "console-test"}
    status, got = call(env, "POST /api/builds/{id}/test-tool", params={"id": "pclaims01"}, body={"key": "docs", "event": event})
    assert status == 200 and got["ok"]
    assert seen == {"fn": f"ToolLambda-{agent}-docs", "key": "docs", "tool": "generate_document", "event": event}
