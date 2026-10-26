"""Rendering for the Document generator (app/tools/docgen/handler.py): PDF, PPTX and DOCX
from JSON data, either into a template (PPTX, DOCX) or into the built-in simple layout.

No AWS here, so the tests render real files from small fixtures. The libraries are pinned
in requirements.txt (python-pptx, python-docx, reportlab; their own dependencies lxml,
Pillow, XlsxWriter, typing_extensions, charset-normalizer) and installed into this folder
by the deploy (deployer/runner.py install_builtin_requirements).

TEMPLATE RULES (docs/WORKFLOW_REFERENCE.md "Document generator" says the same)

  * A field is `{{name}}`, or `{{a.b}}` for a value inside an object; spaces inside the
    braces are allowed. It is replaced by the value in `data`: text as it is, a number or
    true/false as written, a list of plain values joined with ", ", anything else as JSON.
    A field `data` does not have renders as nothing.
  * A list repeats what uses it. `{{rows.name}}`, where `data.rows` is a list, makes:
      - a table row (DOCX or PPTX) that uses it: one row per item;
      - a DOCX paragraph outside a table that uses it: one paragraph per item;
      - a PPTX slide whose text outside its tables uses it: one slide per item.
    Each copy reads `{{rows.<field>}}` from its own item (a plain value is `{{rows.value}}`);
    an empty list removes the row, paragraph or slide. One list per row, paragraph or slide
    (the first one it names); a list inside an item is not repeated again.
  * A paragraph is filled as a whole, so a field Word or PowerPoint split over several
    runs still works; the filled text takes the formatting of the paragraph's first run.
  * PDF has no template format: it always uses the built-in layout.

BUILT-IN LAYOUT (no template): `title` (and `subtitle`) as the heading, then each other
top-level field as a section in the order given: a plain value as a line, an object as
"key: value" lines, a list of plain values as bullets, a list of objects as a table.
"""
from __future__ import annotations

import copy
import io
import json
import re

#: Items one list may repeat a row, paragraph or slide into, and rows one table may have
#: in the built-in layout. Bounds the work one call can ask for.
MAX_ITEMS = 1000
#: Columns the built-in layout shows for a list of objects (the first keys seen).
MAX_COLUMNS = 8
#: Table rows on one built-in PPTX slide before the table continues on the next.
ROWS_PER_SLIDE = 12

FIELD_RE = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_]*(?:\.[A-Za-z0-9_]+)*)\s*\}\}")
#: Characters XML 1.0 refuses (lxml raises on them): C0 controls other than tab, LF, CR.
_XML_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class RenderError(ValueError):
    """The data or the template cannot be rendered; the message says why."""


# --- fields -------------------------------------------------------------------------------

def as_text(value) -> str:
    """A value as the text a field shows."""
    if value is None:
        out = ""
    elif isinstance(value, bool):
        out = "true" if value else "false"
    elif isinstance(value, (str, int, float)):
        out = str(value)
    elif isinstance(value, list) and all(isinstance(v, (str, int, float, bool)) or v is None for v in value):
        out = ", ".join(as_text(v) for v in value)
    else:
        out = json.dumps(value, ensure_ascii=False, default=str)
    return _XML_BAD.sub("", out)


def _lookup(obj, parts: list[str]):
    for p in parts:
        if isinstance(obj, dict) and p in obj:
            obj = obj[p]
        else:
            return None
    return obj


def fill(text: str, data: dict, scope: tuple[str, object] | None = None) -> str:
    """Replace every field in `text`. Inside a repeat, `scope` = (list name, item), and
    `{{<list>.<field>}}` reads the item."""
    def one(m: re.Match) -> str:
        parts = m.group(1).split(".")
        if scope is not None and parts[0] == scope[0] and len(parts) > 1:
            item = scope[1] if isinstance(scope[1], dict) else {"value": scope[1]}
            return as_text(_lookup(item, parts[1:]))
        return as_text(_lookup(data, parts))
    return FIELD_RE.sub(one, text)


def repeated_list(text: str, data: dict) -> str | None:
    """The list a piece of text repeats over: the first `{{<list>.<field>}}` it has whose
    `<list>` is a list in `data`. None when it has none."""
    for m in FIELD_RE.finditer(text):
        parts = m.group(1).split(".")
        if len(parts) > 1 and isinstance(data.get(parts[0]), list):
            return parts[0]
    return None


def _items(data: dict, name: str) -> list:
    items = data.get(name) or []
    if len(items) > MAX_ITEMS:
        raise RenderError(f"data.{name} has {len(items)} items; a template repeats at most {MAX_ITEMS}")
    return items


def _fill_runs(runs: list, data: dict, scope=None) -> None:
    """Fill one paragraph given as its runs: the whole text into the first run, the others
    emptied, so a field split over runs is still found. Untouched without a field."""
    runs = list(runs)
    if not runs:
        return
    full = "".join(r.text for r in runs)
    if "{{" not in full:
        return
    runs[0].text = fill(full, data, scope)
    for r in runs[1:]:
        r.text = ""


# --- DOCX ---------------------------------------------------------------------------------

def _docx_paragraph_text(p_el) -> str:
    from docx.oxml.ns import qn
    return "".join(t.text or "" for t in p_el.iter(qn("w:t")))


def _docx_fill_element(el, data: dict, scope=None) -> None:
    """Fill every paragraph inside an XML element (a row, a cell, a body)."""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    for p_el in list(el.iter(qn("w:p"))):
        _fill_runs(Paragraph(p_el, None).runs, data, scope)


def _docx_repeat(el, data: dict) -> None:
    """Repeat one row or paragraph element once per item of the list it names."""
    name = repeated_list(_docx_paragraph_text(el), data)
    if name is None:
        return
    anchor = el
    for item in _items(data, name):
        clone = copy.deepcopy(el)
        _docx_fill_element(clone, data, (name, item))
        anchor.addnext(clone)
        anchor = clone
    el.getparent().remove(el)


def render_docx(data: dict, template: bytes | None) -> bytes:
    import docx
    from docx.oxml.ns import qn
    if template is None:
        return _builtin_docx(data)
    try:
        doc = docx.Document(io.BytesIO(template))
    except Exception as e:
        raise RenderError(f"the template is not a readable .docx ({type(e).__name__})") from e
    body = doc.element.body
    for tr in list(body.iter(qn("w:tr"))):
        _docx_repeat(tr, data)
    for p_el in [p for p in body.iterchildren(qn("w:p"))]:
        _docx_repeat(p_el, data)
    _docx_fill_element(body, data)
    for section in doc.sections:
        for part in (section.header, section.footer):
            if not part.is_linked_to_previous:
                _docx_fill_element(part._element, data)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# --- PPTX ---------------------------------------------------------------------------------

def _pptx_paragraphs(el, outside_tables: bool = False) -> list:
    from pptx.oxml.ns import qn
    out = []
    for p_el in el.iter(qn("a:p")):
        if outside_tables and any(a.tag == qn("a:tbl") for a in p_el.iterancestors()):
            continue
        out.append(p_el)
    return out


def _pptx_text(p_els) -> str:
    from pptx.oxml.ns import qn
    return "\n".join("".join(t.text or "" for t in p.iter(qn("a:t"))) for p in p_els)


def _pptx_fill_element(el, data: dict, scope=None) -> None:
    from pptx.text.text import _Paragraph
    for p_el in _pptx_paragraphs(el):
        _fill_runs(_Paragraph(p_el, None).runs, data, scope)


def _pptx_duplicate_slide(prs, index: int):
    """A copy of slide `index`, placed right after it: its shapes, background and the
    relationships they use (pictures, links), with the copy's own relationship ids."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn
    src = prs.slides[index]
    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for shape_el in [e for e in tree.iterchildren() if e.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr"))]:
        tree.remove(shape_el)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        rid_map[rid] = (new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref) if rel.is_external
                        else new.part.rels.get_or_add(rel.reltype, rel.target_part))
    for shape_el in src.shapes._spTree.iterchildren():
        if shape_el.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            tree.append(copy.deepcopy(shape_el))
    bg = src._element.cSld.bg
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    r_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for el in new._element.iter():
        for attr, val in list(el.attrib.items()):
            if attr.startswith(r_ns) and val in rid_map:
                el.set(attr, rid_map[val])
    ids = prs.slides._sldIdLst
    moved = ids[-1]
    ids.remove(moved)
    ids.insert(index + 1, moved)
    return new


def _pptx_delete_slide(prs, index: int) -> None:
    ids = prs.slides._sldIdLst
    sld_id = ids[index]
    ids.remove(sld_id)
    prs.part.drop_rel(sld_id.rId)


def render_pptx(data: dict, template: bytes | None) -> bytes:
    import pptx
    from pptx.oxml.ns import qn
    if template is None:
        return _builtin_pptx(data)
    try:
        prs = pptx.Presentation(io.BytesIO(template))
    except Exception as e:
        raise RenderError(f"the template is not a readable .pptx ({type(e).__name__})") from e
    # Slides first: a slide whose own text (outside its tables) names a list.
    i = 0
    while i < len(prs.slides):
        slide = prs.slides[i]
        name = repeated_list(_pptx_text(_pptx_paragraphs(slide.shapes._spTree, outside_tables=True)), data)
        if name is None:
            i += 1
            continue
        items = _items(data, name)
        if not items:
            _pptx_delete_slide(prs, i)
            continue
        copies = [slide] + [_pptx_duplicate_slide(prs, i + n) for n in range(len(items) - 1)]
        for s, item in zip(copies, items, strict=True):
            _pptx_fill_element(s.shapes._spTree, data, (name, item))
        i += len(items)
    for slide in prs.slides:
        tree = slide.shapes._spTree
        for tr in list(tree.iter(qn("a:tr"))):
            name = repeated_list(_pptx_text(_pptx_paragraphs(tr)), data)
            if name is None:
                continue
            anchor = tr
            for item in _items(data, name):
                clone = copy.deepcopy(tr)
                _pptx_fill_element(clone, data, (name, item))
                anchor.addnext(clone)
                anchor = clone
            tr.getparent().remove(tr)
        _pptx_fill_element(tree, data)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# --- the built-in layout ------------------------------------------------------------------

def _label(key: str) -> str:
    return _XML_BAD.sub("", str(key).replace("_", " ").strip().capitalize())


def sections(data: dict) -> tuple[str, str, list[tuple[str, str, object]]]:
    """(title, subtitle, sections) of the built-in layout. Each section is (label, kind,
    content): kind "text" (a string), "pairs" (a list of (key, value)), "bullets" (a list of
    strings) or "table" (header, rows)."""
    title = as_text(data.get("title")) or "Document"
    subtitle = as_text(data.get("subtitle"))
    out: list[tuple[str, str, object]] = []
    for key, value in data.items():
        if key in ("title", "subtitle"):
            continue
        label = _label(key)
        if isinstance(value, dict):
            out.append((label, "pairs", [(_label(k), as_text(v)) for k, v in value.items()]))
        elif isinstance(value, list) and value and all(isinstance(v, dict) for v in value):
            if len(value) > MAX_ITEMS:
                raise RenderError(f"data.{key} has {len(value)} rows; a table holds at most {MAX_ITEMS}")
            header: list[str] = []
            for row in value:
                for k in row:
                    if k not in header and len(header) < MAX_COLUMNS:
                        header.append(k)
            rows = [[as_text(row.get(k)) for k in header] for row in value]
            out.append((label, "table", ([_label(k) for k in header], rows)))
        elif isinstance(value, list):
            if len(value) > MAX_ITEMS:
                raise RenderError(f"data.{key} has {len(value)} items; a list holds at most {MAX_ITEMS}")
            out.append((label, "bullets", [as_text(v) for v in value]))
        else:
            out.append((label, "text", as_text(value)))
    return title, subtitle, out


def render_pdf(data: dict) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    title, subtitle, parts = sections(data)

    def para(text: str, style: str = "BodyText"):
        return Paragraph(escape(text).replace("\n", "<br/>"), styles[style])

    story: list = [para(title, "Title")]
    if subtitle:
        story.append(para(subtitle, "Heading3"))
    for label, kind, content in parts:
        story.append(para(label, "Heading2"))
        if kind == "text":
            story.append(para(content))
        elif kind == "pairs":
            story.extend(para(f"{k}: {v}") for k, v in content)
        elif kind == "bullets":
            story.append(ListFlowable([ListItem(para(b)) for b in content], bulletType="bullet"))
        else:
            header, rows = content
            table = Table([[para(h) for h in header]] + [[para(c) for c in r] for r in rows], repeatRows=1)
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(table)
        story.append(Spacer(1, 4 * mm))
    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=A4, title=title).build(story)
    return out.getvalue()


def _builtin_docx(data: dict) -> bytes:
    import docx
    title, subtitle, parts = sections(data)
    doc = docx.Document()
    doc.add_heading(title, 0)
    if subtitle:
        doc.add_paragraph(subtitle)
    for label, kind, content in parts:
        doc.add_heading(label, 1)
        if kind == "text":
            doc.add_paragraph(content)
        elif kind == "pairs":
            for k, v in content:
                doc.add_paragraph(f"{k}: {v}")
        elif kind == "bullets":
            for b in content:
                doc.add_paragraph(b, style="List Bullet")
        else:
            header, rows = content
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for cell, h in zip(table.rows[0].cells, header, strict=True):
                cell.text = h
            for r in rows:
                for cell, c in zip(table.add_row().cells, r, strict=True):
                    cell.text = c
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _builtin_pptx(data: dict) -> bytes:
    import pptx
    from pptx.util import Inches, Pt
    title, subtitle, parts = sections(data)
    prs = pptx.Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[0])          # Title Slide
    first.shapes.title.text = title
    first.placeholders[1].text = subtitle

    def bullets_slide(heading: str, lines: list[str]) -> None:
        s = prs.slides.add_slide(prs.slide_layouts[1])          # Title and Content
        s.shapes.title.text = heading
        frame = s.placeholders[1].text_frame
        frame.text = lines[0] if lines else ""
        for line in lines[1:]:
            frame.add_paragraph().text = line

    for label, kind, content in parts:
        if kind == "text":
            bullets_slide(label, [content])
        elif kind == "pairs":
            bullets_slide(label, [f"{k}: {v}" for k, v in content])
        elif kind == "bullets":
            bullets_slide(label, list(content))
        else:
            header, rows = content
            for start in range(0, max(len(rows), 1), ROWS_PER_SLIDE):
                chunk = rows[start:start + ROWS_PER_SLIDE]
                s = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
                s.shapes.title.text = label
                shape = s.shapes.add_table(len(chunk) + 1, len(header), Inches(0.5), Inches(1.5),
                                           prs.slide_width - Inches(1), Inches(0.4) * (len(chunk) + 1))
                for c, h in enumerate(header):
                    shape.table.cell(0, c).text = h
                for r, row in enumerate(chunk, 1):
                    for c, value in enumerate(row):
                        cell = shape.table.cell(r, c)
                        cell.text = value
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                run.font.size = Pt(12)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def render(fmt: str, data: dict, template: bytes | None) -> bytes:
    """The file for one call: `fmt` is pdf, pptx or docx; `template` the template's bytes,
    or None for the built-in layout (always None for pdf)."""
    if fmt == "pdf":
        return render_pdf(data)
    if fmt == "pptx":
        return render_pptx(data, template)
    if fmt == "docx":
        return render_docx(data, template)
    raise RenderError(f"unknown format {fmt!r}")
